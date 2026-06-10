from __future__ import annotations

from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.database import SessionLocal, configure_database, init_db
from app.main import app
from app.models.audit_event import AuditEventRecord
from app.models.document import DocumentRecord
from app.models.reference_index import ReferenceIndexRecord


def _pdf_bytes() -> bytes:
    import fitz

    document = fitz.open()
    page = document.new_page()
    page.insert_text((36, 36), "Phase 12B PDF", fontsize=10)
    return document.tobytes()


@pytest.fixture()
def client(tmp_path: Path):
    configure_database(f"sqlite:///{tmp_path / 'phase12b.db'}")
    init_db(drop_existing=True)
    return TestClient(app)


def _create_bundle(client: TestClient) -> str:
    response = client.post("/api/bundles", json={"bundle_number": "PHASE12B-001", "customer_name": "Demo"})
    assert response.status_code == 201, response.text
    return response.json()["id"]


def _upload(client: TestClient, bundle_id: str, document_type: str) -> str:
    response = client.post(
        f"/api/bundles/{bundle_id}/documents",
        data={"document_type": document_type},
        files={"file": ("document.pdf", _pdf_bytes(), "application/pdf")},
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


def test_manual_reference_index_records_manual_source_and_safe_evidence(client: TestClient):
    bundle_id = _create_bundle(client)
    document_id = _upload(client, bundle_id, "VENDOR_INVOICE")

    response = client.patch(
        f"/api/documents/{document_id}/extracted-data",
        headers={"X-Request-ID": "req-phase12b"},
        json={
            "actor": "qa",
            "reason": "field missing",
            "fields": {
                "vendor_invoice_no": "INV-1",
                "vendor_invoice_date": "12-02-2026",
                "vendor_name": "SUPREME COMPUTERS INDIA P LTD",
                "po_reference": "1PTR2526000467",
                "invoice_total": "554,600.00",
            },
        },
    )

    assert response.status_code == 200, response.text
    extracted = response.json()["metadata"]["extracted_data"]
    assert extracted["invoice_total"] == 554600
    refs = response.json()["references"]
    po_ref = next(ref for ref in refs if ref["reference_type"] == "po_reference")
    assert po_ref["source_type"] == "manual_entry"
    assert po_ref["document_type"] == "VENDOR_INVOICE"
    assert po_ref["field_name"] == "po_reference"
    assert po_ref["confidence"] == 1.0
    assert po_ref["evidence_text"] == "Manual entry"
    audit = response.json()["audit_event"]
    assert audit["payload"]["request_id"] == "req-phase12b"
    assert audit["payload"]["reason"] == "field missing"
    assert any(change["field"] == "invoice_total" and change["new_value"] == 554600 for change in audit["payload"]["changes"])


def test_reference_rebuild_deletes_stale_refs_and_deduplicates(client: TestClient):
    bundle_id = _create_bundle(client)
    document_id = _upload(client, bundle_id, "VENDOR_INVOICE")
    first = client.patch(
        f"/api/documents/{document_id}/extracted-data",
        json={
            "reason": "field missing",
            "fields": {
                "vendor_invoice_no": "INV-1",
                "vendor_invoice_date": "12-02-2026",
                "vendor_name": "SUPREME COMPUTERS INDIA P LTD",
                "po_reference": "PO-OLD",
                "customer_ref_no": "PO-OLD",
                "invoice_total": 100,
            },
        },
    )
    assert first.status_code == 200, first.text

    second = client.patch(
        f"/api/documents/{document_id}/extracted-data",
        json={
            "reason": "corrected wrong value",
            "fields": {
                "po_reference": "PO-NEW",
                "customer_ref_no": "",
                "invoice_total": 100,
            },
        },
    )

    assert second.status_code == 200, second.text
    with SessionLocal() as db:
        refs = list(db.scalars(select(ReferenceIndexRecord).where(ReferenceIndexRecord.document_id == document_id)))
    values = [(ref.reference_type, ref.reference_value) for ref in refs]
    assert ("po_reference", "PO-NEW") in values
    assert ("po_reference", "PO-OLD") not in values
    assert len(values) == len(set(values))


def test_manual_patch_does_not_relabel_untouched_extracted_reference_as_manual(client: TestClient):
    bundle_id = _create_bundle(client)
    document_id = _upload(client, bundle_id, "VENDOR_INVOICE")
    with SessionLocal() as db:
        document = db.get(DocumentRecord, document_id)
        document.metadata_record.extracted_data = {
            "vendor_invoice_no": "INV-EXTRACTED",
            "vendor_invoice_date": "12-02-2026",
            "vendor_name": "SUPREME COMPUTERS INDIA P LTD",
            "po_reference": "PO-EXTRACTED",
            "invoice_total": 100,
        }
        document.metadata_record.diagnostics = {}
        db.commit()

    response = client.patch(
        f"/api/documents/{document_id}/extracted-data",
        json={"reason": "corrected wrong value", "fields": {"invoice_total": 125}},
    )

    assert response.status_code == 200, response.text
    refs = response.json()["references"]
    po_ref = next(ref for ref in refs if ref["reference_type"] == "po_reference")
    assert po_ref["source_type"] == "extracted"


def test_manual_patch_requires_reason_for_high_impact_fields(client: TestClient):
    bundle_id = _create_bundle(client)
    document_id = _upload(client, bundle_id, "VENDOR_INVOICE")

    response = client.patch(
        f"/api/documents/{document_id}/extracted-data",
        json={"fields": {"po_reference": "1PTR2526000467"}},
    )

    assert response.status_code == 400
    assert "reason" in response.json()["detail"].lower()


def test_manual_patch_rejects_blank_critical_reference_value(client: TestClient):
    bundle_id = _create_bundle(client)
    document_id = _upload(client, bundle_id, "VENDOR_INVOICE")

    response = client.patch(
        f"/api/documents/{document_id}/extracted-data",
        json={"reason": "corrected wrong value", "fields": {"po_reference": "   "}},
    )

    assert response.status_code == 400
    assert "po_reference" in response.json()["detail"]


def test_manual_patch_and_audit_rollback_together(client: TestClient, monkeypatch: pytest.MonkeyPatch):
    import app.repositories.audit as audit_repository

    bundle_id = _create_bundle(client)
    document_id = _upload(client, bundle_id, "VENDOR_INVOICE")
    original_create = audit_repository.AuditRepository.create

    def fail_create(self, payload):
        raise RuntimeError("audit write failed")

    monkeypatch.setattr(audit_repository.AuditRepository, "create", fail_create)
    response = client.patch(
        f"/api/documents/{document_id}/extracted-data",
        json={
            "reason": "field missing",
            "fields": {
                "vendor_invoice_no": "INV-1",
                "vendor_invoice_date": "12-02-2026",
                "vendor_name": "SUPREME COMPUTERS INDIA P LTD",
                "po_reference": "1PTR2526000467",
                "invoice_total": 100,
            },
        },
    )

    assert response.status_code == 500
    monkeypatch.setattr(audit_repository.AuditRepository, "create", original_create)
    with SessionLocal() as db:
        document = db.get(DocumentRecord, document_id)
        assert document.metadata_record.extracted_data == {}
        assert list(db.scalars(select(AuditEventRecord).where(AuditEventRecord.document_id == document_id))) == []


def test_export_includes_reference_provenance_and_audit_request_id(client: TestClient):
    from io import BytesIO

    from openpyxl import load_workbook

    bundle_id = _create_bundle(client)
    document_id = _upload(client, bundle_id, "VENDOR_INVOICE")
    patch = client.patch(
        f"/api/documents/{document_id}/extracted-data",
        headers={"X-Request-ID": "req-export"},
        json={
            "actor": "qa",
            "reason": "field missing",
            "fields": {
                "vendor_invoice_no": "INV-1",
                "vendor_invoice_date": "12-02-2026",
                "vendor_name": "SUPREME COMPUTERS INDIA P LTD",
                "po_reference": "1PTR2526000467",
                "invoice_total": 554600,
            },
        },
    )
    assert patch.status_code == 200, patch.text

    export = client.get(f"/api/bundles/{bundle_id}/export.xlsx")

    assert export.status_code == 200, export.text
    workbook = load_workbook(BytesIO(export.content))
    refs_sheet = workbook["References"]
    headers = [cell.value for cell in next(refs_sheet.iter_rows(min_row=1, max_row=1))]
    assert "Source Type" in headers
    assert "Evidence Text" in headers
    rows = list(refs_sheet.iter_rows(min_row=2, values_only=True))
    assert any("po_reference" in row and "manual_entry" in row for row in rows)
    audit_headers = [cell.value for cell in next(workbook["Audit Trail"].iter_rows(min_row=1, max_row=1))]
    assert "Request ID" in audit_headers
    audit_rows = list(workbook["Audit Trail"].iter_rows(min_row=2, values_only=True))
    assert any("field missing" in row and "req-export" in row for row in audit_rows)
