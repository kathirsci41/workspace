"""Build 1 live-hardening tests.

Covers the changes made for the full live E2E hardening phase:
- VENDOR_PO is honoured in the PaddleOCR document-type allowlist.
- GLM fallback stays disabled (GLM is not silently routed to).
- DB audit events are persisted for bundle/document/export mutations.
- The export workbook exposes the required Build 1 sheets.
- Invalid (non-PDF) uploads are rejected with a clear, user-visible error.

These tests use fake ``%PDF`` byte payloads and never invoke OCR, so they run
without a GPU. Live OCR behaviour is proven separately by the Playwright E2E
suite and the manual live-run report under ``.runtime/outputs``.
"""
from __future__ import annotations

import dataclasses
import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import settings
from app.database import configure_database, init_db
from app.main import app
from app.services import extraction_service


def _patch_settings(monkeypatch: pytest.MonkeyPatch, **overrides) -> None:
    """Settings is a frozen dataclass; swap the module-level instance the
    extraction-service helpers read for one carrying the overrides."""
    replaced = dataclasses.replace(settings, **overrides)
    monkeypatch.setattr(extraction_service, "settings", replaced)


@pytest.fixture()
def client(tmp_path: Path) -> TestClient:
    configure_database(f"sqlite:///{tmp_path / 'order_assurance_hardening.db'}")
    init_db(drop_existing=True)
    return TestClient(app)


def _create_bundle(client: TestClient, number: str = "E2E-LIVE-TEST-001") -> str:
    response = client.post(
        "/api/bundles",
        json={"bundle_number": number, "customer_name": "TEST CUSTOMER"},
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


def _upload(client: TestClient, bundle_id: str, document_type: str, filename: str, content: bytes) -> "tuple[int, dict]":
    response = client.post(
        f"/api/bundles/{bundle_id}/documents",
        data={"document_type": document_type},
        files={"file": (filename, content, "application/pdf")},
    )
    body = response.json() if response.headers.get("content-type", "").startswith("application/json") else {}
    return response.status_code, body


def _audit_event_types(client: TestClient, bundle_id: str) -> list[str]:
    response = client.get(f"/api/bundles/{bundle_id}/audit-events")
    assert response.status_code == 200, response.text
    return [event["event_type"] for event in response.json()]


# --------------------------------------------------------------------------- #
# Allowlist / GLM-fallback configuration
# --------------------------------------------------------------------------- #
def test_vendor_po_can_be_included_in_paddle_allowlist(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch_settings(
        monkeypatch,
        ocr_paddle_document_types="VENDOR_INVOICE,CUSTOMER_PO,CUSTOMER_INVOICE,VENDOR_PO",
    )
    allowlist = extraction_service._paddle_document_type_allowlist()
    assert "VENDOR_PO" in allowlist
    assert {"VENDOR_INVOICE", "CUSTOMER_PO", "CUSTOMER_INVOICE", "VENDOR_PO"} <= allowlist


def test_vendor_po_routes_to_paddle_when_allowlisted(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch_settings(
        monkeypatch,
        ocr_enabled=True,
        ocr_provider=extraction_service._PADDLE_GPU_PROVIDER_NAME,
        ocr_paddle_document_types="VENDOR_PO,VENDOR_INVOICE",
    )
    assert extraction_service._paddleocr_route_enabled("VENDOR_PO") is True
    # A type left out of the allowlist must not route to paddle.
    assert extraction_service._paddleocr_route_enabled("DELIVERY_CHALLAN") is False


def test_glm_not_routed_when_fallback_disabled(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch_settings(
        monkeypatch,
        ocr_enabled=True,
        ocr_provider=extraction_service._PADDLE_GPU_PROVIDER_NAME,
        ocr_paddle_document_types="VENDOR_INVOICE",
        ocr_paddle_fallback_to_glm=False,
    )
    # DELIVERY_CHALLAN is outside the allowlist; with fallback disabled it must
    # NOT be silently routed to GLM.
    assert extraction_service._glm_route_enabled("DELIVERY_CHALLAN") is False


def test_glm_route_allowed_only_when_fallback_enabled(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch_settings(
        monkeypatch,
        ocr_enabled=True,
        ocr_provider=extraction_service._PADDLE_GPU_PROVIDER_NAME,
        ocr_paddle_document_types="VENDOR_INVOICE",
        ocr_paddle_fallback_to_glm=True,
    )
    assert extraction_service._glm_route_enabled("DELIVERY_CHALLAN") is True


# --------------------------------------------------------------------------- #
# Audit events persisted to the DB
# --------------------------------------------------------------------------- #
def test_bundle_created_audit_event(client: TestClient) -> None:
    bundle_id = _create_bundle(client)
    assert "bundle_created" in _audit_event_types(client, bundle_id)


def test_document_uploaded_audit_event(client: TestClient) -> None:
    bundle_id = _create_bundle(client)
    status_code, body = _upload(client, bundle_id, "CUSTOMER_PO", "po.pdf", b"%PDF-1.4 fake")
    assert status_code == 201, body
    types = _audit_event_types(client, bundle_id)
    assert "document_uploaded" in types
    assert "bundle_created" in types


def test_document_deleted_audit_event(client: TestClient) -> None:
    bundle_id = _create_bundle(client)
    status_code, body = _upload(client, bundle_id, "CUSTOMER_PO", "po.pdf", b"%PDF-1.4 fake")
    assert status_code == 201, body
    document_id = body["id"]
    delete_response = client.delete(f"/api/documents/{document_id}")
    assert delete_response.status_code == 204, delete_response.text
    assert "document_deleted" in _audit_event_types(client, bundle_id)


def test_export_generated_audit_event(client: TestClient) -> None:
    bundle_id = _create_bundle(client)
    _upload(client, bundle_id, "CUSTOMER_PO", "po.pdf", b"%PDF-1.4 fake")
    export_response = client.get(f"/api/bundles/{bundle_id}/export.xlsx")
    assert export_response.status_code == 200, export_response.text
    assert "export_generated" in _audit_event_types(client, bundle_id)


# --------------------------------------------------------------------------- #
# Export workbook structure
# --------------------------------------------------------------------------- #
def test_export_workbook_has_required_sheets(client: TestClient) -> None:
    bundle_id = _create_bundle(client)
    _upload(client, bundle_id, "CUSTOMER_PO", "po.pdf", b"%PDF-1.4 fake")
    export_response = client.get(f"/api/bundles/{bundle_id}/export.xlsx")
    assert export_response.status_code == 200, export_response.text
    workbook = load_workbook(io.BytesIO(export_response.content), read_only=True)
    required = {
        "Executive Dashboard",
        "Business Flow",
        "Financial Findings",
        "Findings",
        "Review Actions",
        "Document Register",
        "Extracted Fields",
        "References",
        "Audit Trail",
        "Technical Diagnostics",
    }
    assert required <= set(workbook.sheetnames), f"missing sheets: {required - set(workbook.sheetnames)}"


# --------------------------------------------------------------------------- #
# Invalid upload handling
# --------------------------------------------------------------------------- #
def test_invalid_non_pdf_upload_rejected(client: TestClient) -> None:
    bundle_id = _create_bundle(client)
    response = client.post(
        f"/api/bundles/{bundle_id}/documents",
        data={"document_type": "CUSTOMER_PO"},
        files={"file": ("notes.txt", b"this is not a pdf", "text/plain")},
    )
    assert response.status_code == 400, response.text
    assert "PDF" in response.json()["detail"]


def test_pdf_extension_with_bad_header_rejected(client: TestClient) -> None:
    bundle_id = _create_bundle(client)
    response = client.post(
        f"/api/bundles/{bundle_id}/documents",
        data={"document_type": "CUSTOMER_PO"},
        files={"file": ("fake.pdf", b"GIF89a not really a pdf", "application/pdf")},
    )
    assert response.status_code == 400, response.text
    assert "valid PDF" in response.json()["detail"]
