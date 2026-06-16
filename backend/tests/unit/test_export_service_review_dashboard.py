from types import SimpleNamespace

from openpyxl import Workbook

from app.services.export_service import (
    _financial_summary,
    _section_header,
    _write_executive_dashboard,
    _write_review_actions,
)


def _flatten_values(ws):
    return [cell.value for row in ws.iter_rows() for cell in row if cell.value is not None]


def _review_items_value(ws):
    for row in ws.iter_rows(values_only=True):
        for idx, value in enumerate(row):
            if value == "Review Items" and idx + 1 < len(row):
                return row[idx + 1]
    raise AssertionError("Review Items KPI was not found")


def _billing_over_po_summary():
    return {
        "issues": [
            {
                "code": "VENDOR_PARTIAL_BILLING_REVIEW_REQUIRED",
                "message": "Vendor bills exceed PO total.",
                "vendor_po_total": 512741.20,
                "vendor_invoice_total": 513376.98,
                "difference": -635.78,
            }
        ],
        "checks": [
            {
                "check_id": "VENDOR_BILLING_COVERAGE",
                "check_name": "Vendor Billing Coverage",
                "result": "REVIEW_REQUIRED",
                "left_value": 512741.20,
                "right_value": 513376.98,
                "message": "Vendor billing coverage evaluated against matching invoice references.",
            }
        ],
        "bundle_status": "REVIEW_REQUIRED",
        "customer_delivery_status": "PASS",
        "vendor_procurement_status": "REVIEW_REQUIRED",
        "recommendation": "Review required before closure.",
        "extracted_summary": {},
    }


def test_section_header_spans_requested_columns_for_dashboard_readability():
    wb = Workbook()
    ws = wb.active

    _section_header(ws, 3, "Key Performance Indicators", start_col=5, end_col=8)

    assert "E3:H3" in {str(rng) for rng in ws.merged_cells.ranges}
    assert ws.row_dimensions[3].height >= 22


def test_executive_dashboard_labels_over_po_difference_as_billing_variance():
    summary = _billing_over_po_summary()
    financial = _financial_summary(summary)
    wb = Workbook()
    ws = wb.active
    bundle = SimpleNamespace(
        bundle_number="Trade-001",
        customer_name="Trade",
        customer_po_no="PO-001",
        so_no="SO-001",
    )

    _write_executive_dashboard(ws, bundle, summary, [], financial)

    values = _flatten_values(ws)
    assert "Billing Variance" in values
    assert "+635.78 over PO" in values
    assert "Difference (PO - Bills)" not in values


def test_review_items_kpi_deduplicates_billing_issue_and_matching_check():
    summary = _billing_over_po_summary()
    summary["checks"] = [
        *summary["checks"],
        {"check_id": "CUSTOMER_PO_INVOICE_ORDER_MATCH", "result": "MISMATCH"},
        {"check_id": "CUSTOMER_PO_INVOICE_AMOUNT_MATCH", "result": "REVIEW_REQUIRED"},
        {"check_id": "VENDOR_PO_INVOICE_REFERENCE_MATCH", "result": "REVIEW_REQUIRED"},
        {"check_id": "INVOICE_DC_SO_MATCH", "result": "REVIEW_REQUIRED"},
    ]
    financial = _financial_summary(summary)
    wb = Workbook()
    ws = wb.active
    bundle = SimpleNamespace(
        bundle_number="Trade-001",
        customer_name="Trade",
        customer_po_no="PO-001",
        so_no="SO-001",
    )

    _write_executive_dashboard(ws, bundle, summary, [], financial)

    assert _review_items_value(ws) == 5


def test_review_actions_deduplicates_repeated_vendor_billing_checks():
    summary = {
        "issues": [],
        "checks": [
            {
                "check_id": "VENDOR_BILLING_COVERAGE",
                "check_name": "Vendor Billing Coverage",
                "result": "REVIEW_REQUIRED",
                "left_value": 512741.20,
                "right_value": 513376.98,
                "message": "Billing mismatch",
            },
            {
                "check_id": "VENDOR_BILLING_COVERAGE",
                "check_name": "Vendor Billing Coverage",
                "result": "REVIEW_REQUIRED",
                "left_value": 512741.20,
                "right_value": 513376.98,
                "message": "Billing mismatch",
            },
        ],
    }
    financial = _financial_summary(summary)
    wb = Workbook()
    ws = wb.active

    _write_review_actions(ws, summary, financial)

    issue_labels = [row[1] for row in ws.iter_rows(min_row=4, values_only=True) if row[1]]
    assert issue_labels.count("Vendor Billing Coverage") == 1
