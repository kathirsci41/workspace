from __future__ import annotations

from datetime import datetime
from io import BytesIO
import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side
)
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.repositories.audit import AuditRepository
from app.repositories.bundles import BundleRepository
from app.repositories.documents import DocumentRepository
from app.repositories.reference_index import ReferenceIndexRepository
from app.services.verification_summary_service import build_verification_summary

# ── Colour palette ─────────────────────────────────────────────────────────
_C_HEADER_DARK  = "1F3864"   # dark navy – title bar
_C_HEADER_MID   = "2F5496"   # medium blue – section headers
_C_HEADER_LIGHT = "D6E4F0"   # pale blue – alt rows / sub-headers
_C_WHITE        = "FFFFFF"
_C_BLACK        = "000000"
_C_GREEN_FILL   = "C6EFCE"
_C_GREEN_FONT   = "375623"
_C_AMBER_FILL   = "FFEB9C"
_C_AMBER_FONT   = "9C5700"
_C_RED_FILL     = "FFC7CE"
_C_RED_FONT     = "9C0006"
_C_GREY_FILL    = "F2F2F2"
_C_GREY_FONT    = "595959"
_C_BLUE_LIGHT   = "BDD7EE"

# ── Canonical document type mapping ────────────────────────────────────────
_CANONICAL: dict[str, str] = {
    "CUSTOMER_PO":    "CUSTOMER_PO",
    "COMPANY_INVOICE":"CUSTOMER_INVOICE",
    "CUSTOMER_INVOICE":"CUSTOMER_INVOICE",
    "COMPANY_DC":     "DELIVERY_CHALLAN",
    "DELIVERY_CHALLAN":"DELIVERY_CHALLAN",
    "COMPANY_PO":     "VENDOR_PO",
    "VENDOR_PO":      "VENDOR_PO",
    "VENDOR_INVOICE": "VENDOR_INVOICE",
    "VENDOR_BILL":    "VENDOR_INVOICE",
    "ERP_ORDER_EXPORT":"ERP_ORDER_EXPORT",
}

# Build 1 business flow order
_BUSINESS_FLOW_ORDER = [
    ("CUSTOMER_PO",    "Customer PO",          "Step 1"),
    ("VENDOR_PO",      "Vendor PO",            "Step 2"),
    ("VENDOR_INVOICE", "Vendor Bills",         "Step 3"),
    ("DELIVERY_CHALLAN","Company DC",          "Step 4"),
    ("CUSTOMER_INVOICE","Company Invoice",     "Step 5"),
]

_VERIFICATION_READY = {"EXTRACTED", "MANUAL_ENTRY"}

# ── Public entry point ──────────────────────────────────────────────────────
def build_bundle_export(db: Session, bundle_id: str) -> bytes:
    bundle    = BundleRepository(db).get(bundle_id)
    summary   = build_verification_summary(db, bundle_id)
    documents = DocumentRepository(db).list_for_bundle(bundle_id)
    references = ReferenceIndexRepository(db).list_for_bundle(bundle_id)
    audit_events = AuditRepository(db).list_for_bundle(bundle_id)

    financial = _financial_summary(summary)
    canonical_groups = _build_canonical_groups(documents)

    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Executive Dashboard"
    _write_executive_dashboard(ws1, bundle, summary, documents, financial)

    _write_business_flow(wb.create_sheet("Business Flow"), canonical_groups, summary, financial)
    _write_financial_findings(wb.create_sheet("Financial Findings"), summary, financial)
    _write_findings(wb.create_sheet("Findings"), summary)
    _write_review_actions(wb.create_sheet("Review Actions"), summary, financial)
    _write_document_register(wb.create_sheet("Document Register"), documents)
    _write_fields(wb.create_sheet("Extracted Fields"), documents)
    _write_references(wb.create_sheet("References"), references)
    _write_audit(wb.create_sheet("Audit Trail"), audit_events)
    _write_technical_diagnostics(wb.create_sheet("Technical Diagnostics"), documents)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


# ── Financial summary helpers ───────────────────────────────────────────────
def _financial_summary(summary: dict[str, Any]) -> dict[str, Any]:
    """
    Extract canonical financial values from the verification summary.
    Uses checks and issues – never raw document totals – to avoid double-counting
    when alias records (COMPANY_PO / VENDOR_PO) both exist in the bundle.
    """
    issues = summary.get("issues", [])
    checks = summary.get("checks", [])

    # Prefer billing issue values (most reliable when present)
    billing_issue = next(
        (i for i in issues if i.get("code") in {
            "VENDOR_PARTIAL_BILLING_REVIEW_REQUIRED",
            "VENDOR_BILLING_REVIEW_REQUIRED",
        }),
        None,
    )
    billing_check = next(
        (c for c in checks if c.get("check_id") == "VENDOR_BILLING_COVERAGE"),
        None,
    )

    vendor_po_total     = _num(billing_issue, "vendor_po_total")     if billing_issue else None
    vendor_invoice_total = _num(billing_issue, "vendor_invoice_total") if billing_issue else None
    difference          = billing_issue.get("difference")            if billing_issue else None

    if vendor_po_total is None and billing_check:
        vendor_po_total = _safe_num(billing_check.get("left_value"))
    if vendor_invoice_total is None and billing_check:
        vendor_invoice_total = _safe_num(billing_check.get("right_value"))
    if difference is None and vendor_po_total is not None and vendor_invoice_total is not None:
        raw_diff = vendor_po_total - vendor_invoice_total
        difference = int(raw_diff) if float(raw_diff).is_integer() else round(raw_diff, 2)

    extracted = summary.get("extracted_summary", {})

    return {
        "vendor_po_total": vendor_po_total,
        "vendor_invoice_total": vendor_invoice_total
            or _safe_num(extracted.get("vendor_total")),
        "difference": difference,
        "has_billing_data": vendor_po_total is not None or vendor_invoice_total is not None,
    }


def _build_canonical_groups(documents) -> dict[str, list]:
    """Group raw documents by canonical type (preserving all records)."""
    groups: dict[str, list] = {}
    for doc in documents:
        ctype = _CANONICAL.get(str(doc.document_type).upper(), str(doc.document_type).upper())
        groups.setdefault(ctype, []).append(doc)
    return groups


# ── Sheet 1: Executive Dashboard ────────────────────────────────────────────
def _write_executive_dashboard(ws, bundle, summary, documents, financial):
    ws.sheet_view.showGridLines = False

    # Title band
    _merge_write(ws, 1, 1, 1, 8,
                 "ORDER ASSURANCE  ·  VERIFICATION REPORT",
                 font=Font(name="Calibri", size=20, bold=True, color=_C_WHITE),
                 fill=PatternFill("solid", fgColor=_C_HEADER_DARK),
                 alignment=Alignment(horizontal="center", vertical="center", wrap_text=True))
    ws.row_dimensions[1].height = 36

    # Bundle meta
    extracted = summary.get("extracted_summary", {})
    bundle_no  = getattr(bundle, "bundle_number", "-") or "-"
    customer   = getattr(bundle, "customer_name", None) or extracted.get("customer_name") or "-"
    po_no      = getattr(bundle, "customer_po_no", None) or extracted.get("customer_po_no") or "-"
    so_no      = getattr(bundle, "so_no", None) or extracted.get("so_no") or "-"
    generated  = datetime.now().strftime("%Y-%m-%d %H:%M")

    _section_header(ws, 3, "Bundle Details")
    meta_rows = [
        ("Bundle Number",    bundle_no),
        ("Customer Name",    customer),
        ("Customer PO No",   po_no),
        ("SO Number",        so_no),
        ("Report Generated", generated),
    ]
    for i, (label, value) in enumerate(meta_rows, start=4):
        _label_value_row(ws, i, label, value)

    # Status section
    bundle_status   = summary.get("bundle_status", "-")
    customer_status = summary.get("customer_delivery_status", "-")
    vendor_status   = summary.get("vendor_procurement_status", "-")
    rec             = summary.get("recommendation") or "-"

    _section_header(ws, 11, "Status")
    _label_value_row(ws, 12, "Bundle Status",            bundle_status,   _status_fill(bundle_status),   _status_font(bundle_status))
    _label_value_row(ws, 13, "Customer Delivery Status", customer_status, _status_fill(customer_status), _status_font(customer_status))
    _label_value_row(ws, 14, "Vendor Procurement Status",vendor_status,   _status_fill(vendor_status),   _status_font(vendor_status))
    _label_value_row(ws, 15, "Recommendation",           rec)
    ws.row_dimensions[15].height = 30

    # KPI cards (cols 5-8)
    total    = len(documents)
    extracted_cnt = sum(1 for d in documents if _meta_status(d) in _VERIFICATION_READY)
    pending  = sum(1 for d in documents if _meta_status(d) not in _VERIFICATION_READY
                   and _doc_status(d) != "EXTRACTION_FAILED")
    failed   = sum(1 for d in documents if _doc_status(d) == "EXTRACTION_FAILED"
                   or _meta_status(d) == "FAILED")
    checks   = summary.get("checks", [])
    review_cnt = _review_item_count(summary)
    mismatch_cnt = sum(1 for c in checks if c.get("result") == "MISMATCH")
    canonical_cnt = len({_CANONICAL.get(str(d.document_type).upper()) for d in documents})

    kpis = [
        ("Uploaded Records",    total,        "info"),
        ("Business Doc Groups", canonical_cnt, "info"),
        ("Extracted",          extracted_cnt, "success" if extracted_cnt == total else "warning"),
        ("Pending",            pending,       "warning" if pending else "success"),
        ("Review Items",       review_cnt,    "warning" if review_cnt else "success"),
        ("Critical Mismatches",mismatch_cnt,  "danger"  if mismatch_cnt else "success"),
    ]
    _section_header_cols(ws, 3, 5, 8, "Key Performance Indicators")
    for i, (label, value, tone) in enumerate(kpis):
        row = 4 + i
        fill_hex = {"success": _C_GREEN_FILL, "warning": _C_AMBER_FILL,
                    "danger": _C_RED_FILL, "info": _C_BLUE_LIGHT}.get(tone, _C_GREY_FILL)
        _kpi_cell(ws, row, 5, label, value, fill_hex)

    # Main finding
    _section_header(ws, 12, "Main Finding", start_col=5, end_col=8)
    main_finding = _main_finding(summary, financial)
    _merge_write(ws, 13, 5, 16, 8, main_finding,
                 font=Font(name="Calibri", size=11, color=_C_BLACK),
                 fill=PatternFill("solid", fgColor=_C_AMBER_FILL
                                  if bundle_status == "REVIEW_REQUIRED" else _C_GREEN_FILL),
                 alignment=Alignment(horizontal="left", vertical="top", wrap_text=True))
    ws.row_dimensions[13].height = 60

    # Financial summary
    _section_header(ws, 18, "Financial Summary")
    fin_rows = []
    if financial["vendor_po_total"] is not None:
        fin_rows.append(("Vendor PO Total",     _fmt_currency(financial["vendor_po_total"])))
    if financial["vendor_invoice_total"] is not None:
        fin_rows.append(("Vendor Bills Total",  _fmt_currency(financial["vendor_invoice_total"])))
    if financial["difference"] is not None:
        fin_rows.append(("Billing Variance", _billing_variance_text(financial["difference"])))
    if not fin_rows:
        fin_rows.append(("Note", "Financial data not yet available — extraction pending."))
    for i, (label, value) in enumerate(fin_rows, start=19):
        _label_value_row(ws, i, label, value)

    # Column widths
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 32
    ws.column_dimensions["D"].width = 4
    ws.column_dimensions["E"].width = 24
    ws.column_dimensions["F"].width = 20
    ws.column_dimensions["G"].width = 4
    ws.column_dimensions["H"].width = 24
    ws.freeze_panes = "B3"


# ── Sheet 2: Business Flow ──────────────────────────────────────────────────
def _write_business_flow(ws, canonical_groups, summary, financial):
    ws.sheet_view.showGridLines = False
    _sheet_title(ws, 1, "Business Document Flow")

    headers = ["Step", "Business Document", "Status", "Key Reference", "Amount", "Uploaded", "Extracted", "Notes"]
    _table_header(ws, 3, headers)

    checks = summary.get("checks", [])
    extracted = summary.get("extracted_summary", {})
    issues    = summary.get("issues", [])

    billing_check = next((c for c in checks if c.get("check_id") == "VENDOR_BILLING_COVERAGE"), None)

    row = 4
    for canonical_type, label, step in _BUSINESS_FLOW_ORDER:
        docs = canonical_groups.get(canonical_type, [])
        if not docs and canonical_type == "VENDOR_INVOICE":
            docs = canonical_groups.get("VENDOR_BILL", [])

        uploaded  = len(docs)
        extr_cnt  = sum(1 for d in docs if _meta_status(d) in _VERIFICATION_READY)
        status    = _group_status(docs, canonical_type, issues, checks)

        # Multiple vendor invoices get separate rows
        if canonical_type == "VENDOR_INVOICE" and docs:
            vendor_refs  = _dedup_list(extracted.get("vendor_invoice_numbers", []))
            for i, doc in enumerate(docs):
                meta = doc.metadata_record
                exdata = meta.extracted_data if meta else {}
                inv_no  = (exdata.get("vendor_invoice_no") or
                           exdata.get("invoice_no") or
                           (vendor_refs[i] if i < len(vendor_refs) else "-"))
                amount  = _doc_amount(doc, "invoice_total", "net_amount", "total_amount")
                row_status = status if extr_cnt == uploaded else ("PENDING" if _meta_status(doc) not in _VERIFICATION_READY else "EXTRACTED")
                fill = _status_fill(row_status)
                note = "Matched to Vendor PO" if (billing_check or issues) else ""
                _table_row(ws, row, [
                    step if i == 0 else "",
                    f"{label} {i+1}" if len(docs) > 1 else label,
                    row_status,
                    str(inv_no) if inv_no else "-",
                    _fmt_currency(amount) if amount is not None else "-",
                    "1" if i == 0 else "",
                    "1" if _meta_status(doc) in _VERIFICATION_READY else "0",
                    note,
                ], fill)
                row += 1
            continue

        # Single row for other types
        ref    = _canonical_key_ref(canonical_type, docs, extracted)
        amount = _canonical_amount(canonical_type, docs, financial)
        fill   = _status_fill(status)
        notes  = _flow_notes(canonical_type, issues, checks)
        _table_row(ws, row, [
            step, label, status, ref,
            _fmt_currency(amount) if amount is not None else "-",
            str(uploaded), str(extr_cnt), notes,
        ], fill)
        row += 1

    _auto_filter(ws, 3, len(headers))
    _set_col_widths(ws, [6, 22, 18, 24, 16, 10, 10, 40])
    ws.freeze_panes = "A4"


# ── Sheet 3: Financial Findings ──────────────────────────────────────────────
def _write_financial_findings(ws, summary, financial):
    ws.sheet_view.showGridLines = False
    _sheet_title(ws, 1, "Financial Findings")

    headers = ["Area", "Expected Amount", "Actual Amount", "Difference", "Result", "Explanation"]
    _table_header(ws, 3, headers)

    checks  = summary.get("checks", [])
    issues  = summary.get("issues", [])

    row = 4
    # Vendor billing comparison
    po_total  = financial["vendor_po_total"]
    inv_total = financial["vendor_invoice_total"]
    diff      = financial["difference"]

    if po_total is not None or inv_total is not None:
        if diff is not None and isinstance(diff, (int, float)):
            if abs(diff) <= 2:
                result = "PASS"
                explanation = "Vendor bills match Vendor PO total within rounding tolerance."
            elif diff > 0:
                result = "REVIEW_REQUIRED"
                explanation = f"Vendor bills are ₹{_fmt_currency(abs(diff))} below the Vendor PO total. Confirm partial billing or upload remaining invoice."
            else:
                result = "REVIEW_REQUIRED"
                explanation = f"Vendor bills exceed Vendor PO total by ₹{_fmt_currency(abs(diff))}. Review with procurement/finance before closure."
        else:
            result = "REVIEW_REQUIRED"
            explanation = "Financial comparison requires manual review."
        fill = _status_fill(result)
        _table_row(ws, row, [
            "Vendor Billing Coverage",
            _fmt_currency(po_total) if po_total is not None else "-",
            _fmt_currency(inv_total) if inv_total is not None else "-",
            (f"+{_fmt_currency(abs(diff))}" if isinstance(diff, (int, float)) and diff < 0
             else _fmt_currency(abs(diff)) if isinstance(diff, (int, float)) else "-"),
            result,
            explanation,
        ], fill)
        row += 1

    # Customer amount check (if present)
    cust_check = next((c for c in checks if c.get("check_id") == "CUSTOMER_PO_INVOICE_AMOUNT_MATCH"), None)
    if cust_check:
        result = cust_check.get("result", "-")
        fill = _status_fill(result)
        _table_row(ws, row, [
            "Customer PO vs Invoice Amount",
            _fmt_currency_safe(cust_check.get("left_value")),
            _fmt_currency_safe(cust_check.get("right_value")),
            "-",
            result,
            cust_check.get("message", "-"),
        ], fill)
        row += 1

    if row == 4:
        _table_row(ws, row, ["No financial data", "—", "—", "—", "N/A",
                              "Extraction has not completed yet."], None)

    _auto_filter(ws, 3, len(headers))
    _set_col_widths(ws, [28, 16, 16, 14, 16, 50])
    ws.freeze_panes = "A4"


# ── Sheet 4: Findings ────────────────────────────────────────────────────────
def _write_findings(ws, summary):
    ws.sheet_view.showGridLines = False
    _sheet_title(ws, 1, "Verification Findings")

    headers = ["Check Area", "Left Document", "Left Value", "Right Document", "Right Value", "Result", "Finding", "Action Required"]
    _table_header(ws, 3, headers)

    row = 4
    for check in summary.get("checks", []):
        result     = check.get("result", "-")
        fill       = _status_fill(result)
        left_type  = _nice_doc_type(check.get("left_document_type"))
        right_type = _nice_doc_type(check.get("right_document_type"))
        action     = ("Review and correct extracted values." if result != "PASS"
                      else "No action required.")
        _table_row(ws, row, [
            _nice_check_name(check.get("check_name") or check.get("check_id")),
            left_type,
            _fmt_value(check.get("left_value")),
            right_type,
            _fmt_value(check.get("right_value")),
            result,
            check.get("message", "-"),
            action,
        ], fill)
        row += 1

    if row == 4:
        _table_row(ws, row, ["No checks available", "—", "—", "—", "—", "—",
                              "Extraction has not completed yet.", "—"], None)

    _auto_filter(ws, 3, len(headers))
    _set_col_widths(ws, [28, 18, 18, 18, 18, 14, 44, 36])
    ws.freeze_panes = "A4"


# ── Sheet 5: Review Actions ──────────────────────────────────────────────────
def _write_review_actions(ws, summary, financial):
    ws.sheet_view.showGridLines = False
    _sheet_title(ws, 1, "Review Actions")

    headers = ["Priority", "Issue", "Evidence", "Owner Team", "Suggested Action", "Status"]
    _table_header(ws, 3, headers)

    issues = summary.get("issues", [])
    seen_action_keys: set[str] = set()
    row    = 4

    for issue in issues:
        action_key = _review_action_key_for_issue(issue)
        if action_key in seen_action_keys:
            continue
        seen_action_keys.add(action_key)
        code    = str(issue.get("code") or "")
        message = str(issue.get("message") or "")
        priority, team, action, evidence = _issue_action_data(code, message, issue, financial)
        fill = PatternFill("solid", fgColor=_C_AMBER_FILL if priority == "Medium"
                           else (_C_RED_FILL if priority == "High" else _C_GREY_FILL))
        _table_row(ws, row, [priority, _friendly_issue(code, message), evidence, team, action, "Open"], fill)
        row += 1

    for check in summary.get("checks", []):
        if not _is_review_check(check):
            continue
        action_key = _review_action_key_for_check(check)
        if action_key in seen_action_keys:
            continue
        seen_action_keys.add(action_key)
        check_id = str(check.get("check_id") or "")
        priority, team, action, evidence = _check_action_data(check_id, check, financial)
        fill = PatternFill("solid", fgColor=_C_AMBER_FILL if priority == "Medium"
                           else (_C_RED_FILL if priority == "High" else _C_GREY_FILL))
        _table_row(ws, row, [
            priority,
            _nice_check_name(check_id),
            evidence,
            team,
            action,
            "Open",
        ], fill)
        row += 1

    if row == 4:
        _merge_write(ws, 4, 1, 4, 6, "No review actions required. All checks passed.",
                     font=Font(name="Calibri", size=11, color=_C_GREEN_FONT),
                     fill=PatternFill("solid", fgColor=_C_GREEN_FILL),
                     alignment=Alignment(horizontal="center", vertical="center"))

    _auto_filter(ws, 3, len(headers))
    _set_col_widths(ws, [10, 36, 44, 20, 48, 10])
    ws.freeze_panes = "A4"


# ── Sheet 6: Document Register ───────────────────────────────────────────────
def _write_document_register(ws, documents):
    ws.sheet_view.showGridLines = False
    _sheet_title(ws, 1, "Document Register")

    headers = [
        "Stored Document Type", "Canonical Document Type", "Display Label",
        "Filename", "Document Status", "Metadata Status",
        "Extraction Route", "Is Alias Record", "Verification Ready", "Uploaded At",
    ]
    _table_header(ws, 3, headers)

    primary_seen: dict[str, bool] = {}
    for row_idx, doc in enumerate(documents, start=4):
        raw_type   = str(doc.document_type).upper()
        canon_type = _CANONICAL.get(raw_type, raw_type)
        is_alias   = raw_type != canon_type
        is_primary = not primary_seen.get(canon_type, False)
        if is_primary and not is_alias:
            primary_seen[canon_type] = True

        meta       = doc.metadata_record
        diag       = meta.diagnostics if meta else {}
        ms         = meta.status if meta else "-"
        ready      = "Yes" if ms in _VERIFICATION_READY else "No"
        fill       = (PatternFill("solid", fgColor=_C_GREY_FILL) if is_alias else None)

        _table_row(ws, row_idx, [
            raw_type,
            canon_type,
            _nice_doc_type(raw_type),
            doc.filename,
            doc.status,
            ms,
            diag.get("extraction_route") or "-",
            "Yes" if is_alias else "No",
            ready,
            _excel_value(doc.created_at),
        ], fill)

    _auto_filter(ws, 3, len(headers))
    _set_col_widths(ws, [20, 20, 20, 36, 16, 16, 16, 14, 14, 20])
    ws.freeze_panes = "A4"


# ── Sheet 7: Extracted Fields ────────────────────────────────────────────────
def _write_fields(ws, documents):
    ws.sheet_view.showGridLines = False
    _sheet_title(ws, 1, "Extracted Fields")

    headers = [
        "Document Filename", "Canonical Type", "Stored Type",
        "Field Name", "Extracted Value", "Confidence", "Source",
        "Page", "Evidence Text",
    ]
    _table_header(ws, 3, headers)

    row = 4
    for doc in documents:
        meta = doc.metadata_record
        if not meta:
            continue
        raw_type   = str(doc.document_type).upper()
        canon_type = _CANONICAL.get(raw_type, raw_type)
        source     = meta.extracted_data.get("extraction_source")
        fm         = meta.diagnostics.get("field_metadata") or {}
        fl         = meta.diagnostics.get("field_locations") or {}
        for field, value in meta.extracted_data.items():
            if str(field).startswith("raw_"):
                continue
            if isinstance(value, (list, dict)):
                continue
            det  = fm.get(field, {})
            loc  = fl.get(field, {})
            _table_row(ws, row, [
                doc.filename,
                canon_type,
                raw_type,
                field,
                value,
                det.get("confidence", meta.diagnostics.get("parser_confidence")),
                det.get("source") or source,
                loc.get("page"),
                loc.get("evidence_text") or det.get("evidence_text"),
            ], None)
            row += 1

    _auto_filter(ws, 3, len(headers))
    _set_col_widths(ws, [36, 18, 18, 24, 28, 12, 14, 8, 40])
    ws.freeze_panes = "A4"


# ── Sheet 8: References ──────────────────────────────────────────────────────
def _write_references(ws, references):
    ws.sheet_view.showGridLines = False
    _sheet_title(ws, 1, "Cross-Document References")

    headers = [
        "Reference Type", "Source Document Type", "Reference Value",
        "Source Type", "Confidence", "Evidence Text", "Created At",
    ]
    _table_header(ws, 3, headers)

    for row_idx, ref in enumerate(references, start=4):
        _table_row(ws, row_idx, [
            ref.reference_type,
            ref.document_type,
            ref.reference_value,
            ref.source_type,
            ref.confidence,
            ref.evidence_text,
            _excel_value(ref.created_at),
        ], None)

    if not references:
        _table_row(ws, 4, ["No cross-document references indexed yet.", "", "", "", "", "", ""], None)

    _auto_filter(ws, 3, len(headers))
    _set_col_widths(ws, [24, 20, 24, 16, 12, 36, 20])
    ws.freeze_panes = "A4"


# ── Sheet 9: Audit Trail ─────────────────────────────────────────────────────
def _write_audit(ws, events):
    ws.sheet_view.showGridLines = False
    _sheet_title(ws, 1, "Audit Trail")

    headers = ["Timestamp", "Document Type", "Action", "Field", "Old Value", "New Value", "Source / Actor", "Reason", "Request ID"]
    _table_header(ws, 3, headers)

    row = 4
    for event in events:
        payload = event.payload or {}
        changes = payload.get("changes") or []
        if not changes:
            _table_row(ws, row, [
                _excel_value(event.created_at),
                payload.get("document_type"),
                event.event_type,
                None, None, None,
                payload.get("source") or event.actor,
                payload.get("reason"),
                payload.get("request_id"),
            ], None)
            row += 1
            continue
        for change in changes:
            _table_row(ws, row, [
                _excel_value(event.created_at),
                payload.get("document_type"),
                event.event_type,
                change.get("field"),
                change.get("old_value"),
                change.get("new_value"),
                payload.get("source") or event.actor,
                payload.get("reason"),
                payload.get("request_id"),
            ], None)
            row += 1

    if row == 4:
        _merge_write(ws, 4, 1, 4, 8,
                     "No audit events recorded for this bundle.",
                     font=Font(name="Calibri", size=11, color=_C_GREY_FONT),
                     fill=PatternFill("solid", fgColor=_C_GREY_FILL),
                     alignment=Alignment(horizontal="center", vertical="center"))

    _auto_filter(ws, 3, len(headers))
    _set_col_widths(ws, [20, 20, 28, 20, 20, 20, 20, 36, 20])
    ws.freeze_panes = "A4"


# ── Sheet 10: Technical Diagnostics ─────────────────────────────────────────
def _write_technical_diagnostics(ws, documents):
    ws.sheet_view.showGridLines = False
    _sheet_title(ws, 1, "Technical Diagnostics")

    headers = [
        "Document Filename", "Stored Type", "Canonical Type",
        "Extraction Route", "OCR Provider", "Digital Text Used",
        "Text Length", "Parser Confidence", "Last Error", "Notes",
    ]
    _table_header(ws, 3, headers)

    for row_idx, doc in enumerate(documents, start=4):
        raw_type   = str(doc.document_type).upper()
        canon_type = _CANONICAL.get(raw_type, raw_type)
        meta       = doc.metadata_record
        diag       = meta.diagnostics if meta else {}
        _table_row(ws, row_idx, [
            doc.filename,
            raw_type,
            canon_type,
            diag.get("extraction_route"),
            diag.get("ocr_provider"),
            diag.get("digital_text_used"),
            diag.get("digital_text_length"),
            diag.get("parser_confidence"),
            meta.last_error if meta else None,
            diag.get("failure_reason"),
        ], None)

    _auto_filter(ws, 3, len(headers))
    _set_col_widths(ws, [36, 18, 18, 16, 14, 14, 12, 16, 28, 28])
    ws.freeze_panes = "A4"


# ── Style helpers ────────────────────────────────────────────────────────────
def _sheet_title(ws, row: int, title: str) -> None:
    cell = ws.cell(row=row, column=1, value=title)
    cell.font       = Font(name="Calibri", size=14, bold=True, color=_C_WHITE)
    cell.fill       = PatternFill("solid", fgColor=_C_HEADER_MID)
    cell.alignment  = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 22


def _section_header(ws, row: int, title: str, start_col: int = 1, end_col: int = 3) -> None:
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(name="Calibri", size=11, bold=True, color=_C_WHITE)
        cell.fill = PatternFill("solid", fgColor=_C_HEADER_MID)
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    if end_col > start_col:
        ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
    ws.cell(row=row, column=start_col, value=title)
    ws.row_dimensions[row].height = 22


def _section_header_cols(ws, row: int, start_col: int, end_col: int, title: str) -> None:
    _section_header(ws, row, title, start_col, end_col)


def _table_header(ws, row: int, headers: list[str]) -> None:
    for col, label in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.font      = Font(name="Calibri", size=10, bold=True, color=_C_WHITE)
        cell.fill      = PatternFill("solid", fgColor=_C_HEADER_DARK)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = _thin_border()
    ws.row_dimensions[row].height = 18


def _table_row(ws, row: int, values: list, fill: PatternFill | None) -> None:
    alt_fill = PatternFill("solid", fgColor="EBF3FB") if row % 2 == 0 else None
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font      = Font(name="Calibri", size=10)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border    = _thin_border()
        if fill:
            cell.fill = fill
        elif alt_fill:
            cell.fill = alt_fill
    ws.row_dimensions[row].height = 15


def _label_value_row(ws, row: int, label: str, value, fill=None, font_color=None):
    lc = ws.cell(row=row, column=2, value=label)
    lc.font      = Font(name="Calibri", size=10, bold=True, color=_C_BLACK)
    lc.alignment = Alignment(vertical="center", indent=1)
    lc.border    = _thin_border()
    vc = ws.cell(row=row, column=3, value=value)
    vc.font      = Font(name="Calibri", size=10,
                        color=font_color if font_color else _C_BLACK)
    vc.alignment = Alignment(vertical="center", wrap_text=True)
    vc.border    = _thin_border()
    if fill:
        vc.fill = fill
    ws.row_dimensions[row].height = 16


def _merge_write(ws, min_row, min_col, max_row, max_col, value,
                 font=None, fill=None, alignment=None):
    ws.merge_cells(
        start_row=min_row, start_column=min_col,
        end_row=max_row,   end_column=max_col,
    )
    cell            = ws.cell(row=min_row, column=min_col, value=value)
    if font:        cell.font      = font
    if fill:        cell.fill      = fill
    if alignment:   cell.alignment = alignment


def _kpi_cell(ws, row: int, col: int, label: str, value, fill_hex: str):
    lc = ws.cell(row=row, column=col, value=label)
    lc.font      = Font(name="Calibri", size=9, bold=True, color=_C_BLACK)
    lc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    lc.fill      = PatternFill("solid", fgColor=fill_hex)
    lc.border    = _thin_border()
    vc = ws.cell(row=row, column=col + 1, value=value)
    vc.font      = Font(name="Calibri", size=11, bold=True, color=_C_BLACK)
    vc.alignment = Alignment(horizontal="center", vertical="center")
    vc.fill      = PatternFill("solid", fgColor=fill_hex)
    vc.border    = _thin_border()
    ws.row_dimensions[row].height = 18


def _auto_filter(ws, header_row: int, col_count: int):
    last_col = get_column_letter(col_count)
    ws.auto_filter.ref = f"A{header_row}:{last_col}{header_row}"


def _set_col_widths(ws, widths: list[int]):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _thin_border() -> Border:
    s = Side(style="thin", color="D0D0D0")
    return Border(left=s, right=s, top=s, bottom=s)


def _status_fill(status: str | None) -> PatternFill:
    s = str(status or "").upper()
    if s in ("PASS", "OK", "EXTRACTED", "VERIFIED"):
        return PatternFill("solid", fgColor=_C_GREEN_FILL)
    if s in ("MISMATCH", "EXTRACTION_FAILED", "FAILED", "MISSING_DOCUMENTS"):
        return PatternFill("solid", fgColor=_C_RED_FILL)
    if s in ("PENDING", "UPLOADING", "UPLOADED", "EXTRACTING"):
        return PatternFill("solid", fgColor=_C_GREY_FILL)
    return PatternFill("solid", fgColor=_C_AMBER_FILL)


def _status_font(status: str | None) -> str:
    s = str(status or "").upper()
    if s in ("PASS", "OK", "EXTRACTED", "VERIFIED"):
        return _C_GREEN_FONT
    if s in ("MISMATCH", "EXTRACTION_FAILED", "FAILED", "MISSING_DOCUMENTS"):
        return _C_RED_FONT
    return _C_AMBER_FONT


# ── Domain helpers ────────────────────────────────────────────────────────────
def _main_finding(summary: dict, financial: dict) -> str:
    issues  = summary.get("issues", [])
    checks  = summary.get("checks", [])
    b_status = str(summary.get("bundle_status", "")).upper()
    c_status = str(summary.get("customer_delivery_status", "")).upper()
    v_status = str(summary.get("vendor_procurement_status", "")).upper()

    _REVIEW_STATUSES = {"REVIEW_REQUIRED", "MISMATCH", "MISSING", "PARTIAL_PASS"}
    has_bad_checks = any(
        str(c.get("result", "")).upper() not in {"PASS", "OK"} for c in checks
    )
    has_bad_status = (
        b_status in _REVIEW_STATUSES or
        c_status in _REVIEW_STATUSES or
        v_status in _REVIEW_STATUSES
    )

    def _billing_text() -> str | None:
        billing = next(
            (c for c in checks
             if c.get("check_id") == "VENDOR_BILLING_COVERAGE"
             and str(c.get("result", "")).upper() not in {"PASS", "OK"}),
            None,
        )
        if not billing:
            return None
        po   = financial.get("vendor_po_total")
        inv  = financial.get("vendor_invoice_total")
        diff = financial.get("difference")
        if po is None or inv is None or diff is None:
            return None
        direction = "above" if isinstance(diff, (int, float)) and diff < 0 else "below"
        return (
            f"Vendor bills ({_fmt_currency(inv)}) were matched to the Vendor PO "
            f"({_fmt_currency(po)}), but the billing total is ₹{_fmt_currency(abs(diff))} "
            f"{direction} the PO value. "
            "Review with procurement/finance before closure."
        )

    if not issues:
        if has_bad_checks or has_bad_status:
            text = _billing_text()
            if text:
                return text
            bad_names = [
                _nice_check_name(c.get("check_name") or c.get("check_id"))
                for c in checks
                if str(c.get("result", "")).upper() not in {"PASS", "OK"}
            ]
            if bad_names:
                return (
                    f"Review required: {', '.join(bad_names[:3])}. "
                    "Confirm document bundle before closure."
                )
            return "Review items require attention before closure."
        return "All verification checks passed. The order bundle references and totals are consistent."

    first = issues[0]
    code  = str(first.get("code") or "")
    msg   = str(first.get("message") or "")
    if "PARTIAL_BILLING" in code or "BILLING" in code:
        text = _billing_text()
        if text:
            return text
    return msg or summary.get("recommendation") or "Review items require attention before closure."


def _canonical_key_ref(canonical_type: str, docs: list, extracted: dict) -> str:
    if not docs:
        return "-"
    doc  = next((d for d in docs if _meta_status(d) in _VERIFICATION_READY), docs[0])
    meta = doc.metadata_record
    exd  = meta.extracted_data if meta else {}
    ref_field_map = {
        "CUSTOMER_PO":    ("customer_po_no",),
        "VENDOR_PO":      ("vendor_po_no",),
        "CUSTOMER_INVOICE":("invoice_no",),
        "DELIVERY_CHALLAN":("dc_no",),
        "VENDOR_INVOICE": ("vendor_invoice_no", "invoice_no"),
    }
    for field in ref_field_map.get(canonical_type, ()):
        val = exd.get(field)
        if val:
            return str(val)
    return "-"


def _canonical_amount(canonical_type: str, docs: list, financial: dict) -> float | None:
    if canonical_type == "VENDOR_PO":
        return financial.get("vendor_po_total")
    if canonical_type == "VENDOR_INVOICE":
        return financial.get("vendor_invoice_total")
    if not docs:
        return None
    doc  = next((d for d in docs if _meta_status(d) in _VERIFICATION_READY), docs[0])
    return _doc_amount(doc, "grand_total", "net_amount", "invoice_total", "total_amount")


def _doc_amount(doc, *fields) -> float | None:
    meta = doc.metadata_record if doc else None
    exd  = meta.extracted_data if meta else {}
    for f in fields:
        val = exd.get(f)
        if val is not None:
            try:
                n = float(str(val).replace(",", ""))
                return int(n) if n.is_integer() else n
            except (TypeError, ValueError):
                pass
    return None


def _group_status(docs: list, canonical_type: str, issues: list, checks: list) -> str:
    if not docs:
        return "MISSING"
    ready = [d for d in docs if _meta_status(d) in _VERIFICATION_READY]
    if not ready:
        return "PENDING"
    check_ids = {
        "CUSTOMER_PO":     ["CUSTOMER_PO_INVOICE_ORDER_MATCH", "CUSTOMER_PO_INVOICE_AMOUNT_MATCH"],
        "VENDOR_PO":       ["VENDOR_BILLING_COVERAGE", "VENDOR_PO_INVOICE_REFERENCE_MATCH"],
        "VENDOR_INVOICE":  ["VENDOR_BILLING_COVERAGE", "VENDOR_NAME_MATCH"],
        "DELIVERY_CHALLAN":["INVOICE_DC_CUSTOMER_ORDER_MATCH", "INVOICE_DC_SO_MATCH"],
        "CUSTOMER_INVOICE":["INVOICE_DC_CUSTOMER_ORDER_MATCH"],
    }.get(canonical_type, [])
    relevant = [c for c in checks if c.get("check_id") in check_ids]
    if any(c.get("result") == "MISMATCH" for c in relevant):
        return "MISMATCH"
    if any(c.get("result") == "REVIEW_REQUIRED" for c in relevant):
        return "REVIEW_REQUIRED"
    if relevant:
        return "PASS"
    return "EXTRACTED"


def _flow_notes(canonical_type: str, issues: list, checks: list) -> str:
    if canonical_type == "VENDOR_PO":
        billing = next((c for c in checks if c.get("check_id") == "VENDOR_BILLING_COVERAGE"), None)
        if billing:
            return f"Billing check: {billing.get('result', '-')}"
    return ""


def _review_item_count(summary: dict) -> int:
    seen: set[str] = set()
    for issue in summary.get("issues", []):
        seen.add(_review_action_key_for_issue(issue))
    for check in summary.get("checks", []):
        if _is_review_check(check):
            seen.add(_review_action_key_for_check(check))
    return len(seen)


def _is_review_check(check: dict) -> bool:
    return str(check.get("result", "")).upper() not in {"PASS", "OK"}


def _review_action_key_for_issue(issue: dict) -> str:
    code = str(issue.get("code") or "").upper()
    message = str(issue.get("message") or "").upper()
    if "BILLING" in code or "BILLING" in message:
        return "CHECK:VENDOR_BILLING_COVERAGE"
    return "|".join(
        [
            "ISSUE",
            code,
            str(issue.get("document_type") or "").upper(),
            str(issue.get("document_id") or ""),
            str(issue.get("message") or ""),
        ]
    )


def _review_action_key_for_check(check: dict) -> str:
    check_id = str(check.get("check_id") or "").upper()
    if check_id == "VENDOR_BILLING_COVERAGE":
        return "CHECK:VENDOR_BILLING_COVERAGE"
    return "|".join(
        [
            "CHECK",
            check_id,
            str(check.get("left_document_id") or ""),
            str(check.get("right_document_id") or ""),
            str(check.get("left_value") or ""),
            str(check.get("right_value") or ""),
        ]
    )


def _issue_action_data(code: str, message: str, issue: dict, financial: dict):
    text = f"{code} {message}".upper()
    if "PARTIAL_BILLING" in text or "BILLING" in text:
        po   = financial.get("vendor_po_total")
        inv  = financial.get("vendor_invoice_total")
        diff = financial.get("difference")
        ev   = (f"Vendor PO {_fmt_currency(po)} vs Vendor Bills {_fmt_currency(inv)}"
                if po is not None and inv is not None else message)
        action = ("Confirm price, tax, rounding, or request vendor clarification."
                  if diff is not None and isinstance(diff, (int, float)) and abs(diff) < 5000
                  else "Confirm whether this is partial billing. Upload remaining invoice if applicable.")
        return "Medium", "Procurement / Finance", action, ev
    if "MISSING" in text:
        doc_type = issue.get("document_type") or "Document"
        return "High", "Operations", f"Upload the missing {_nice_doc_type(doc_type)}.", message
    if "MISMATCH" in text:
        return "High", "Finance / Procurement", "Verify extracted values against the source document.", message
    if "UNMATCHED" in text or "REFERENCE" in text:
        return "Medium", "Procurement", "Confirm the Vendor PO reference on the vendor invoice.", message
    return "Low", "Review Team", "Review the flagged item and confirm or correct extracted values.", message


def _check_action_data(check_id: str, check: dict, financial: dict):
    """Map a failing check_id to (priority, team, action, evidence) for Review Actions."""
    cid = check_id.upper()
    result = str(check.get("result", "")).upper()
    msg = check.get("message") or "-"
    if "BILLING_COVERAGE" in cid:
        po   = financial.get("vendor_po_total")
        inv  = financial.get("vendor_invoice_total")
        diff = financial.get("difference")
        ev = (f"Vendor PO {_fmt_currency(po)} vs Vendor Bills {_fmt_currency(inv)}"
              if po is not None and inv is not None else msg)
        action = ("Confirm price, tax, rounding, or request vendor clarification."
                  if diff is not None and isinstance(diff, (int, float)) and abs(diff) < 5000
                  else "Confirm billing completeness. Upload remaining invoice if partial billing.")
        return "Medium", "Procurement / Finance", action, ev
    if "MISSING" in cid or result == "MISSING":
        return "High", "Operations", "Upload the missing document.", msg
    if "MISMATCH" in cid or result == "MISMATCH":
        return "High", "Finance / Procurement", "Verify extracted values against the source document.", msg
    if "REFERENCE" in cid:
        return "Medium", "Procurement", "Confirm the document reference matches.", msg
    return "Medium", "Review Team", "Review the flagged check and confirm extracted values.", msg


def _nice_doc_type(dtype: str | None) -> str:
    if not dtype:
        return "-"
    labels = {
        "CUSTOMER_PO":     "Customer PO",
        "COMPANY_INVOICE": "Company Invoice",
        "CUSTOMER_INVOICE":"Customer Invoice",
        "COMPANY_DC":      "Company DC",
        "DELIVERY_CHALLAN":"Delivery Challan",
        "COMPANY_PO":      "Company PO",
        "VENDOR_PO":       "Vendor PO",
        "VENDOR_INVOICE":  "Vendor Invoice",
        "VENDOR_BILL":     "Vendor Bill",
    }
    return labels.get(str(dtype).upper(), str(dtype))


def _nice_check_name(name: str | None) -> str:
    if not name:
        return "-"
    return (str(name)
            .replace("_", " ")
            .replace("INVOICE", "Invoice")
            .replace("VENDOR", "Vendor")
            .replace("CUSTOMER", "Customer")
            .title())


def _friendly_issue(code: str, message: str) -> str:
    code_up = code.upper()
    if "PARTIAL_BILLING" in code_up:
        return "Vendor billing total does not match Vendor PO"
    if "VENDOR_BILLING" in code_up:
        return "Vendor billing mismatch"
    if "VENDOR_PO_MISSING" in code_up:
        return "Vendor PO is missing"
    if "VENDOR_INVOICE_MISSING" in code_up:
        return "Vendor invoice is missing"
    if "CUSTOMER_PO_MISSING" in code_up:
        return "Customer PO is missing or unreadable"
    if "MISSING" in code_up:
        return "Required document is missing"
    if "UNMATCHED" in code_up:
        return "Vendor bill reference does not match any Vendor PO"
    return message[:80] if message else code


def _dedup_list(lst: list) -> list:
    seen = set()
    result = []
    for item in lst:
        if item not in seen:
            seen.add(item)
            result.append(item)
    return result


def _meta_status(doc) -> str:
    meta = doc.metadata_record if doc else None
    return meta.status if meta else ""


def _doc_status(doc) -> str:
    return str(doc.status) if doc else ""


def _fmt_currency(value) -> str:
    if value is None:
        return "-"
    try:
        n = float(str(value).replace(",", ""))
        return f"{n:,.2f}"
    except (TypeError, ValueError):
        return str(value)


def _billing_variance_text(value) -> str:
    try:
        n = float(str(value).replace(",", ""))
    except (TypeError, ValueError):
        return str(value)
    amount = _fmt_currency(abs(n))
    if abs(n) <= 2:
        return f"{amount} within tolerance"
    if n < 0:
        return f"+{amount} over PO"
    return f"{amount} under PO"


def _fmt_currency_safe(value) -> str:
    if value is None:
        return "-"
    try:
        n = float(str(value).replace(",", ""))
        return f"{n:,.2f}"
    except (TypeError, ValueError):
        return str(value)


def _fmt_value(value) -> str:
    if value is None:
        return "-"
    if isinstance(value, float):
        return f"{value:,.2f}" if not value.is_integer() else f"{int(value):,}"
    if isinstance(value, int):
        return f"{value:,}"
    return str(value)


def _num(d: dict, key: str) -> float | None:
    return _safe_num(d.get(key))


def _safe_num(value) -> float | None:
    if value is None:
        return None
    try:
        n = float(str(value).replace(",", ""))
        return int(n) if n.is_integer() else n
    except (TypeError, ValueError):
        return None


def _excel_value(value) -> Any:
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    return value
