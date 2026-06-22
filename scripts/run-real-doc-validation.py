from __future__ import annotations

import argparse
import csv
import json
import os
import re
import subprocess
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
import requests


API_DEFAULT = "http://127.0.0.1:8100/api"


@dataclass(frozen=True)
class DocumentInput:
    document_type: str
    path: Path


def main() -> int:
    parser = argparse.ArgumentParser(description="Run real-document validation through the public Order Assurance API.")
    parser.add_argument("--api-base", default=API_DEFAULT)
    parser.add_argument("--run-root", required=True, type=Path)
    parser.add_argument("--bundle-name", required=True)
    parser.add_argument("--expected", type=Path)
    parser.add_argument("--doc", action="append", default=[], help="Document mapping in DOCUMENT_TYPE=path format. Repeat for each PDF.")
    parser.add_argument("--start-backend", action="store_true", help="Start Uvicorn for this validation run and stop it before exit.")
    parser.add_argument("--backend-port", default=8100, type=int)
    args = parser.parse_args()

    api_base = (args.api_base if args.api_base != API_DEFAULT else f"http://127.0.0.1:{args.backend_port}/api").rstrip("/")
    run_root = args.run_root.resolve()
    artifacts = run_root / "artifacts"
    downloads = run_root / "downloads"
    artifacts.mkdir(parents=True, exist_ok=True)
    downloads.mkdir(parents=True, exist_ok=True)

    documents = [_parse_doc_arg(value) for value in args.doc]
    if not documents:
        raise SystemExit("At least one --doc DOCUMENT_TYPE=path argument is required.")

    expected = _read_json(args.expected) if args.expected else None
    _write_json(
        artifacts / "00-input.json",
        {
            "api_base": api_base,
            "run_root": str(run_root),
            "bundle_name": args.bundle_name,
            "expected": str(args.expected.resolve()) if args.expected else None,
            "documents": [{"document_type": doc.document_type, "path": str(doc.path)} for doc in documents],
        },
    )

    backend = _start_backend(run_root, args.backend_port) if args.start_backend else None
    try:
        bundle = _create_bundle(api_base, args.bundle_name)
        bundle_id = bundle["id"]
        _write_json(artifacts / "01-created-bundle.json", bundle)

        uploaded = []
        for doc in documents:
            uploaded_doc = _upload_document(api_base, bundle_id, doc)
            uploaded.append(uploaded_doc)
        _write_json(artifacts / "02-uploaded-documents.json", uploaded)

        extraction_results = []
        for index, document in enumerate(uploaded, start=1):
            result = _post_json(api_base, f"/documents/{document['id']}/extract")
            extraction_results.append(result)
            doc_type = str(document.get("document_type") or f"document-{index}")
            _write_json(artifacts / f"03-extraction-{index:02d}-{_slug(doc_type)}.json", result)

        final_documents = _get_json(api_base, f"/bundles/{bundle_id}/documents")
        verification = _get_json(api_base, f"/bundles/{bundle_id}/verification-summary")
        derived_issues = _derive_issues(verification)
        audit_events = _get_json(api_base, f"/bundles/{bundle_id}/audit-events")
        document_api_responses = [_get_json(api_base, f"/documents/{doc['id']}") for doc in final_documents]

        _write_json(artifacts / "09-documents-final.json", final_documents)
        _write_json(artifacts / "10-verification-summary.json", verification)
        _write_json(artifacts / "11-derived-issues.json", derived_issues)
        _write_json(artifacts / "12-audit-events.json", audit_events)
        _write_json(artifacts / "15-document-api-responses.json", document_api_responses)
        _write_json(artifacts / "16-raw-extracted-fields.json", _raw_fields(final_documents))

        workbook_path = _download_export(api_base, bundle_id, args.bundle_name, downloads)
        workbook_inspection = _inspect_workbook(workbook_path)
        _write_json(artifacts / "14-workbook-inspection.json", workbook_inspection)

        baseline_rows: list[dict[str, Any]] = []
        if expected:
            baseline_rows = _compare_expected(expected, final_documents)
            _write_csv(artifacts / "17-baseline-comparison.csv", baseline_rows)
            _write_json(artifacts / "17-baseline-comparison.json", baseline_rows)

        validation_result = _validation_result(
            bundle,
            final_documents,
            verification,
            derived_issues,
            audit_events,
            workbook_path,
            workbook_inspection,
            baseline_rows,
            expected,
        )
        _write_json(artifacts / "18-validation-result.json", validation_result)
        print(json.dumps(validation_result, indent=2, ensure_ascii=False))
        return 0
    finally:
        if backend:
            _stop_backend(backend)


def _parse_doc_arg(value: str) -> DocumentInput:
    if "=" not in value:
        raise SystemExit(f"Invalid --doc value: {value!r}. Expected DOCUMENT_TYPE=path.")
    document_type, raw_path = value.split("=", 1)
    path = Path(raw_path).expanduser().resolve()
    if not path.is_file():
        raise SystemExit(f"Document path does not exist: {path}")
    return DocumentInput(document_type=document_type.strip().upper(), path=path)


def _start_backend(run_root: Path, port: int) -> subprocess.Popen:
    project_root = Path(__file__).resolve().parents[1]
    backend_root = project_root / "backend"
    env = os.environ.copy()
    env.update(
        {
            "DATABASE_URL": "sqlite:///" + str((run_root / "order_assurance.db").resolve()).replace("\\", "/"),
            "STORAGE_DIR": str((run_root / "storage" / "documents").resolve()),
            "APP_ENV": "development",
            "ENABLE_DEV_TOOLS": "false",
            "BACKEND_PORT": str(port),
        }
    )
    stdout = (run_root / "backend.validation.out.log").open("ab")
    stderr = (run_root / "backend.validation.err.log").open("ab")
    process = subprocess.Popen(
        [sys.executable, "-m", "uvicorn", "app.main:app", "--host", "127.0.0.1", "--port", str(port)],
        cwd=backend_root,
        env=env,
        stdout=stdout,
        stderr=stderr,
    )
    process._oa_stdout = stdout  # type: ignore[attr-defined]
    process._oa_stderr = stderr  # type: ignore[attr-defined]
    _wait_for_health(port, process)
    return process


def _wait_for_health(port: int, process: subprocess.Popen, timeout_seconds: int = 90) -> None:
    deadline = time.time() + timeout_seconds
    url = f"http://127.0.0.1:{port}/api/bundles"
    last_error = None
    while time.time() < deadline:
        if process.poll() is not None:
            raise RuntimeError(f"Backend exited before health check passed. Exit code: {process.returncode}")
        try:
            response = requests.get(url, timeout=2)
            if response.ok:
                return
            last_error = f"HTTP {response.status_code}"
        except Exception as exc:  # noqa: BLE001 - diagnostic loop
            last_error = str(exc)
        time.sleep(1)
    raise RuntimeError(f"Backend did not become healthy at {url}: {last_error}")


def _stop_backend(process: subprocess.Popen) -> None:
    if process.poll() is None:
        process.terminate()
        try:
            process.wait(timeout=15)
        except subprocess.TimeoutExpired:
            process.kill()
            process.wait(timeout=15)
    for attr in ("_oa_stdout", "_oa_stderr"):
        handle = getattr(process, attr, None)
        if handle:
            handle.close()


def _create_bundle(api_base: str, bundle_name: str) -> dict[str, Any]:
    response = requests.post(f"{api_base}/bundles", json={"bundle_number": bundle_name}, timeout=30)
    _raise_for_status(response)
    return response.json()


def _upload_document(api_base: str, bundle_id: str, doc: DocumentInput) -> dict[str, Any]:
    with doc.path.open("rb") as handle:
        response = requests.post(
            f"{api_base}/bundles/{bundle_id}/documents",
            data={"document_type": doc.document_type},
            files={"file": (doc.path.name, handle, "application/pdf")},
            timeout=120,
        )
    _raise_for_status(response)
    return response.json()


def _post_json(api_base: str, path: str) -> dict[str, Any]:
    response = requests.post(f"{api_base}{path}", timeout=1500)
    _raise_for_status(response)
    return response.json()


def _get_json(api_base: str, path: str) -> Any:
    response = requests.get(f"{api_base}{path}", timeout=120)
    _raise_for_status(response)
    return response.json()


def _download_export(api_base: str, bundle_id: str, bundle_name: str, downloads: Path) -> Path:
    response = requests.get(f"{api_base}/bundles/{bundle_id}/export.xlsx", timeout=120)
    _raise_for_status(response)
    filename = _content_disposition_filename(response.headers.get("Content-Disposition")) or f"{bundle_name}-verification-report.xlsx"
    target = downloads / _safe_filename(filename)
    target.write_bytes(response.content)
    return target


def _raise_for_status(response: requests.Response) -> None:
    if response.ok:
        return
    body = response.text[:1000]
    raise requests.HTTPError(f"{response.request.method} {response.url} failed {response.status_code}: {body}", response=response)


def _content_disposition_filename(value: str | None) -> str | None:
    if not value:
        return None
    match = re.search(r'filename="?([^";]+)"?', value, flags=re.I)
    return match.group(1) if match else None


def _safe_filename(value: str) -> str:
    return re.sub(r'[<>:"/\\|?*]+', "-", value).strip() or "verification-report.xlsx"


def _derive_issues(summary: dict[str, Any]) -> list[dict[str, Any]]:
    issues = [dict(issue, source="summary_issue") for issue in summary.get("issues") or []]
    existing = {(issue.get("code"), issue.get("document_id")) for issue in issues}
    for check in summary.get("checks") or []:
        result = str(check.get("result") or "").upper()
        if result not in {"MISMATCH", "REVIEW_REQUIRED", "MISSING_DOCUMENTS", "BLOCKED"}:
            continue
        code = check.get("check_id")
        key = (code, check.get("right_document_id") or check.get("left_document_id"))
        if key in existing:
            continue
        issues.append(
            {
                "source": "check",
                "code": code,
                "message": check.get("message"),
                "document_type": check.get("right_document_type") or check.get("left_document_type"),
                "document_id": check.get("right_document_id") or check.get("left_document_id"),
                "result": result,
                "expected": check.get("left_value"),
                "found": check.get("right_value"),
            }
        )
    return issues


def _raw_fields(documents: list[dict[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for document in documents:
        doc_type = document.get("document_type") or document.get("id")
        result.setdefault(doc_type, []).append(
            {
                "document_id": document.get("id"),
                "filename": document.get("filename"),
                "status": document.get("status"),
                "metadata_status": (document.get("metadata") or {}).get("status"),
                "extracted_data": document.get("extracted_data") or (document.get("metadata") or {}).get("extracted_data") or {},
                "diagnostics": (document.get("metadata") or {}).get("diagnostics") or {},
                "last_error": document.get("last_error") or (document.get("metadata") or {}).get("last_error"),
            }
        )
    return result


def _compare_expected(expected: dict[str, Any], documents: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_type = {str(document.get("document_type")).upper(): document for document in documents}
    rows: list[dict[str, Any]] = []
    for doc_type, doc_expected in (expected.get("documents") or {}).items():
        document = by_type.get(str(doc_type).upper())
        extracted = _extracted_data(document) if document else {}
        for field, expected_value in (doc_expected.get("expected_fields") or {}).items():
            extracted_value = extracted.get(field)
            passed = _values_match(expected_value, extracted_value)
            rows.append(
                {
                    "document_type": doc_type,
                    "field": field,
                    "expected": expected_value,
                    "extracted": extracted_value,
                    "pass": passed,
                    "failure_type": "" if passed else _default_failure_type(extracted_value),
                    "suggested_fix_area": "" if passed else "extraction pipeline",
                }
            )
    return rows


def _extracted_data(document: dict[str, Any] | None) -> dict[str, Any]:
    if not document:
        return {}
    return document.get("extracted_data") or (document.get("metadata") or {}).get("extracted_data") or {}


def _values_match(expected: Any, extracted: Any) -> bool:
    if expected in (None, ""):
        return extracted in (None, "")
    expected_number = _as_number(expected)
    extracted_number = _as_number(extracted)
    if expected_number is not None and extracted_number is not None:
        return abs(expected_number - extracted_number) <= 0.01
    return _norm_text(expected) == _norm_text(extracted)


def _as_number(value: Any) -> float | None:
    if isinstance(value, bool) or value in (None, ""):
        return None
    if isinstance(value, int | float):
        return float(value)
    text = re.sub(r"[^0-9.\-]", "", str(value))
    if not re.fullmatch(r"-?\d+(?:\.\d+)?", text):
        return None
    return float(text)


def _norm_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip(" .,\t\r\n")).casefold()


def _default_failure_type(extracted: Any) -> str:
    return "extraction missed value" if extracted in (None, "") else "extraction wrong value"


def _inspect_workbook(path: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(path, data_only=True)
    result: dict[str, Any] = {
        "path": str(path),
        "sheets": workbook.sheetnames,
        "summary": {},
        "cell_values": [],
    }
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                if value is not None:
                    result["cell_values"].append(value)

    if "Verification Summary" in workbook.sheetnames:
        sheet = workbook["Verification Summary"]
        for key, value in sheet.iter_rows(min_row=1, max_col=2, values_only=True):
            if key is not None:
                result["summary"][str(key)] = value
    elif "Executive Dashboard" in workbook.sheetnames:
        sheet = workbook["Executive Dashboard"]
        for key, value in sheet.iter_rows(min_col=2, max_col=3, values_only=True):
            if key is not None:
                result["summary"][str(key)] = value
    return result


def _validation_result(
    bundle: dict[str, Any],
    documents: list[dict[str, Any]],
    verification: dict[str, Any],
    derived_issues: list[dict[str, Any]],
    audit_events: list[dict[str, Any]],
    workbook_path: Path,
    workbook_inspection: dict[str, Any],
    baseline_rows: list[dict[str, Any]],
    expected: dict[str, Any] | None,
) -> dict[str, Any]:
    baseline_fail = sum(1 for row in baseline_rows if not row.get("pass"))
    failed_statuses = [
        {
            "document_type": document.get("document_type"),
            "document_id": document.get("id"),
            "document_status": document.get("status"),
            "metadata_status": (document.get("metadata") or {}).get("status"),
            "last_error": document.get("last_error") or (document.get("metadata") or {}).get("last_error"),
        }
        for document in documents
        if document.get("status") == "EXTRACTION_FAILED" or (document.get("metadata") or {}).get("status") == "FAILED"
    ]
    export_fail = 0
    expected_values = ((expected or {}).get("export") or {}).get("expected_values") or []
    summary_values = list((workbook_inspection.get("summary") or {}).values()) + list(workbook_inspection.get("cell_values") or [])
    for expected_value in expected_values:
        if not any(_values_match(expected_value, value) for value in summary_values):
            export_fail += 1
    return {
        "bundle_id": bundle.get("id"),
        "bundle_number": bundle.get("bundle_number"),
        "baseline_rows": len(baseline_rows),
        "baseline_pass": len(baseline_rows) - baseline_fail,
        "baseline_fail": baseline_fail,
        "export_fail": export_fail,
        "failed_document_status_count": len(failed_statuses),
        "failed_document_statuses": failed_statuses,
        "bundle_status": verification.get("bundle_status"),
        "customer_delivery_status": verification.get("customer_delivery_status"),
        "vendor_procurement_status": verification.get("vendor_procurement_status"),
        "check_count": len(verification.get("checks") or []),
        "issue_count": len(verification.get("issues") or []),
        "derived_issue_count": len(derived_issues),
        "audit_event_count": len(audit_events),
        "xlsx_path": str(workbook_path),
        "xlsx_summary": workbook_inspection.get("summary"),
    }


def _read_json(path: Path | None) -> Any:
    if not path:
        return None
    return json.loads(path.read_text(encoding="utf-8"))


def _write_json(path: Path, data: Any) -> None:
    payload = json.dumps(data, indent=2, ensure_ascii=False, default=str)
    try:
        path.write_text(payload, encoding="utf-8")
    except PermissionError:
        if path.exists() and path.is_file():
            path.unlink()
        path.write_text(payload, encoding="utf-8")


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    fieldnames = ["document_type", "field", "expected", "extracted", "pass", "failure_type", "suggested_fix_area"]
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def _slug(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "-", value).strip("-").lower() or "document"


if __name__ == "__main__":
    raise SystemExit(main())
